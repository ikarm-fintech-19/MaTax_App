#!/usr/bin/env python3
"""
build_pptx.py — Compile all 16 SVG slides into a PPTX.

Strategy:
  1. Rasterize each SVG → PNG at 2× (2560×1440) via cairosvg or svglib+reportlab
  2. Insert each PNG as a full-slide image in a 16:9 PPTX (10in × 5.625in)
  3. Export to exports/ with a timestamp

Dependencies (all pre-installed): cairosvg OR svglib/reportlab, python-pptx, Pillow
"""

import os
import sys
import glob
import datetime
import tempfile
from pathlib import Path

# ── Paths ──────────────────────────────────────────────────────────────────
BASE    = Path(__file__).parent
SVG_DIR = BASE / "svg_final"
OUT_DIR = BASE / "exports"
OUT_DIR.mkdir(exist_ok=True)

timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
OUT_PPTX  = OUT_DIR / f"MaTax_Thesis_Defense_{timestamp}.pptx"

# ── Collect slides in order ────────────────────────────────────────────────
svg_files = sorted(SVG_DIR.glob("*.svg"))
if not svg_files:
    print(f"ERROR: No SVG files found in {SVG_DIR}")
    sys.exit(1)

print(f"Found {len(svg_files)} slides:")
for f in svg_files:
    print(f"  {f.name}")

# ── Select rasterizer ──────────────────────────────────────────────────────
def rasterize_cairosvg(svg_path: Path, png_path: Path, scale: float = 2.0):
    import cairosvg
    cairosvg.svg2png(
        url=str(svg_path),
        write_to=str(png_path),
        scale=scale,
    )

def rasterize_svglib(svg_path: Path, png_path: Path, scale: float = 2.0):
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPM
    drawing = svg2rlg(str(svg_path))
    if drawing is None:
        raise RuntimeError(f"svglib could not parse {svg_path}")
    # scale up
    drawing.width  *= scale
    drawing.height *= scale
    drawing.transform = (scale, 0, 0, scale, 0, 0)
    renderPM.drawToFile(drawing, str(png_path), fmt="PNG")

# Try cairosvg first (higher fidelity), fall back to svglib
try:
    import cairosvg
    rasterize = rasterize_cairosvg
    print("\nUsing rasterizer: cairosvg")
except ImportError:
    try:
        import svglib
        rasterize = rasterize_svglib
        print("\nUsing rasterizer: svglib")
    except ImportError:
        print("ERROR: Neither cairosvg nor svglib is installed.")
        sys.exit(1)

# ── Build PPTX ─────────────────────────────────────────────────────────────
from pptx import Presentation
from pptx.util import Inches, Pt

# 16:9 slide size — 10in × 5.625in
SLIDE_W = Inches(10)
SLIDE_H = Inches(5.625)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

blank_layout = prs.slide_layouts[6]  # Blank layout

with tempfile.TemporaryDirectory() as tmpdir:
    for i, svg_path in enumerate(svg_files, start=1):
        png_path = Path(tmpdir) / f"slide_{i:02d}.png"
        print(f"\n[{i:02d}/{len(svg_files)}] Rasterizing {svg_path.name} …", end=" ", flush=True)

        try:
            rasterize(svg_path, png_path, scale=2.0)
            print("OK")
        except Exception as e:
            print(f"FAILED: {e}")
            print("  Skipping this slide.")
            continue

        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(png_path),
            left=0, top=0,
            width=SLIDE_W, height=SLIDE_H,
        )
        print(f"  → Added slide {i}")

prs.save(str(OUT_PPTX))
print(f"\n✅  Saved: {OUT_PPTX}")
print(f"   Slides: {len(prs.slides)}")
